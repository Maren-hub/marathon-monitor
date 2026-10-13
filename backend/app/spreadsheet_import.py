"""Read tabular device data without modifying or executing the uploaded workbook."""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import datetime
from pathlib import PurePath
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pydantic import ValidationError

from .pipeline_schemas import EnvironmentReading, GisSegmentReading, PipelineIngestRequest, UavReading, WearableReading
from .simulation import SEGMENT_DEFINITIONS


MAX_UPLOAD_BYTES = 6 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 80
MAX_ERRORS = 30
SENSITIVE_HEADERS = {"姓名", "手机号", "手机号码", "身份证", "身份证号", "电话号码"}
KNOWN_SEGMENTS = {item["id"] for item in SEGMENT_DEFINITIONS}

MODELS = {
    "wearable": WearableReading,
    "uav": UavReading,
    "environment": EnvironmentReading,
    "gis_segments": GisSegmentReading,
}
SHEET_NAMES = {
    "wearable": {"wearable", "手环", "运动员", "运动员数据", "体征"},
    "uav": {"uav", "drone", "无人机", "无人机数据"},
    "environment": {"environment", "weather", "环境", "环境数据", "天气"},
    "gis_segments": {"gis", "gissegments", "赛段", "赛段gis", "道路", "路段"},
}
ALIASES = {
    "timestamp": ["timestamp", "采集时间", "时间", "记录时间"],
    "source_id": ["source_id", "数据源编号", "设备编号", "传感器编号"],
    "athlete_id": ["athlete_id", "运动员编号", "选手编号"],
    "segment_id": ["segment_id", "赛段编号", "所属赛段", "赛段"],
    "distance_km": ["distance_km", "已完成距离", "已完成距离km", "距离km"],
    "speed_m_s": ["speed_m_s", "当前速度", "当前速度m/s", "运动员速度"],
    "heart_rate_bpm": ["heart_rate_bpm", "心率", "心率bpm"],
    "spo2_pct": ["spo2_pct", "血氧", "血氧饱和度", "血氧饱和度%"],
    "athlete_status": ["athlete_status", "运动员状态", "选手状态"],
    "person_count": ["person_count", "人数", "人员数量", "识别人数"],
    "observed_area_m2": ["observed_area_m2", "观测面积", "观测面积m2", "观测面积㎡"],
    "avg_speed_m_s": ["avg_speed_m_s", "平均速度", "平均速度m/s"],
    "suspected_fall_count": ["suspected_fall_count", "疑似跌倒人数", "疑似跌倒数"],
    "abnormal_motion_count": ["abnormal_motion_count", "异常运动人数", "异常运动数"],
    "temperature_c": ["temperature_c", "温度", "气温", "气温℃"],
    "humidity_pct": ["humidity_pct", "湿度", "相对湿度", "湿度%"],
    "wind_speed_m_s": ["wind_speed_m_s", "风速", "风速m/s"],
    "rainfall_mm": ["rainfall_mm", "降雨量", "降雨量mm"],
    "start_km": ["start_km", "起点里程", "起点里程km"],
    "end_km": ["end_km", "终点里程", "终点里程km"],
    "road_width_m": ["road_width_m", "道路宽度", "道路宽度m"],
    "mean_slope_pct": ["mean_slope_pct", "平均坡度", "平均坡度%"],
    "max_slope_pct": ["max_slope_pct", "最大坡度", "最大坡度%"],
    "curve_level": ["curve_level", "弯道等级"],
    "emergency_access": ["emergency_access", "救援通达性", "应急通行"],
    "nearest_medical_distance_km": ["nearest_medical_distance_km", "最近医疗点距离", "最近医疗点距离km"],
    "data_quality": ["data_quality", "数据质量"],
    "data_status": ["data_status", "数据性质", "数据状态"],
}


def _normalize(value: object) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("_", "").replace("（", "(").replace("）", ")")


FIELD_BY_HEADER = {_normalize(alias): field for field, aliases in ALIASES.items() for alias in aliases}
SENTINELS = {
    "wearable": {"athlete_id", "heart_rate_bpm", "spo2_pct", "distance_km"},
    "uav": {"person_count", "observed_area_m2", "avg_speed_m_s"},
    "environment": {"temperature_c", "humidity_pct", "wind_speed_m_s"},
    "gis_segments": {"road_width_m", "mean_slope_pct", "emergency_access"},
}


class SpreadsheetImportError(Exception):
    def __init__(self, message: str, issues: list[str] | None = None):
        super().__init__(message)
        self.issues = issues or []


def _kind_from_name(sheet_name: str) -> str | None:
    name = _normalize(sheet_name)
    return next((kind for kind, names in SHEET_NAMES.items() if name in names), None)


def _header_fields(cells: tuple) -> dict[int, str]:
    return {index: FIELD_BY_HEADER[_normalize(value)] for index, value in enumerate(cells)
            if _normalize(value) in FIELD_BY_HEADER}


def _detect_kind(fields: set[str], name_hint: str | None) -> str | None:
    scores = {kind: len(fields & sentinels) for kind, sentinels in SENTINELS.items()}
    top = max(scores.values())
    inferred = next((kind for kind, score in scores.items() if score == top), None) if top >= 1 else None
    if name_hint and inferred and name_hint != inferred:
        return None
    return name_hint or inferred


def _clean_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, str):
        return value.strip()
    return value


def _validate_row(kind: str, values: tuple, fields: dict[int, str], default_status: str,
                  location: str, issues: list[str]) -> dict | None:
    row = {field: _clean_value(values[index]) for index, field in fields.items()
           if index < len(values) and values[index] is not None and str(values[index]).strip() != ""}
    if not row:
        return None
    row.setdefault("data_status", default_status)
    if kind != "gis_segments" and isinstance(row.get("timestamp"), (int, float)):
        issues.append(f"{location}：采集时间是纯数字，请在 Excel 中设为日期时间或填 2026-10-10 09:30:00")
        return None
    for identifier in ("source_id", "athlete_id", "segment_id"):
        if identifier in row:
            row[identifier] = str(row[identifier]).strip()
    if row.get("segment_id") and row["segment_id"] not in KNOWN_SEGMENTS:
        issues.append(f"{location}：赛段编号 {row['segment_id']} 不在当前平台赛段 S1–S5 中")
        return None
    try:
        reading = MODELS[kind].model_validate(row)
        if kind == "gis_segments" and reading.end_km <= reading.start_km:
            issues.append(f"{location}：终点里程必须大于起点里程")
            return None
        payload = reading.model_dump(mode="json")
        if kind != "gis_segments" and reading.timestamp.tzinfo is not None:
            payload["timestamp"] = reading.timestamp.astimezone(ZoneInfo("Asia/Shanghai")).replace(tzinfo=None).isoformat(timespec="seconds")
        return payload
    except ValidationError as error:
        for detail in error.errors()[:3]:
            field = str(detail["loc"][0]) if detail["loc"] else "数据"
            issues.append(f"{location}：{field} {detail['msg']}")
        return None


def _read_table(name: str, rows, result: dict[str, list[dict]], issues: list[str],
                report: list[dict], default_status: str, total_rows: list[int]) -> None:
    iterator = iter(enumerate(rows, start=1))
    scanned = [next(iterator, None) for _ in range(10)]
    scanned = [item for item in scanned if item is not None]
    hint = _kind_from_name(name)
    candidate = None
    for row_number, cells in scanned:
        fields = _header_fields(cells)
        kind = _detect_kind(set(fields.values()), hint)
        if kind and len(fields) >= 2:
            candidate = (row_number, fields, kind)
            break
    if candidate is None:
        if hint:
            issues.append(f"工作表“{name}”：前10行未找到可识别的表头")
        return
    header_row, fields, kind = candidate
    if any(_normalize(value) in SENSITIVE_HEADERS for value in scanned[header_row - 1][1]):
        issues.append(f"工作表“{name}”：含姓名、手机号或身份证字段，请先脱敏后上传")
        return
    if len(set(fields.values())) != len(fields):
        issues.append(f"工作表“{name}”：有重复含义的列名，请只保留一列")
        return
    accepted = 0
    def consume(row_number: int, cells: tuple) -> None:
        nonlocal accepted
        if not any(value is not None and str(value).strip() for value in cells):
            return
        total_rows[0] += 1
        if total_rows[0] > MAX_ROWS:
            raise SpreadsheetImportError(f"文件超过 {MAX_ROWS} 行数据上限，请分批导入")
        if len(issues) >= MAX_ERRORS:
            return
        parsed = _validate_row(kind, cells, fields, default_status, f"{name} 第{row_number}行", issues)
        if parsed is not None:
            result[kind].append(parsed)
            accepted += 1
    for row_number, cells in scanned:
        if row_number > header_row:
            consume(row_number, cells)
    for row_number, cells in iterator:
        consume(row_number, cells)
    report.append({"sheet": name, "kind": kind, "rows": accepted})


def parse_uploaded_table(filename: str, data: bytes, default_status: str) -> tuple[PipelineIngestRequest, list[dict]]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise SpreadsheetImportError("文件太大，请上传不超过 6 MB 的表格")
    if default_status not in {"模拟", "公开", "实测"}:
        raise SpreadsheetImportError("请选择正确的数据性质：模拟、公开或实测")
    extension = PurePath(filename).suffix.lower()
    result: dict[str, list[dict]] = {key: [] for key in MODELS}
    issues: list[str] = []
    report: list[dict] = []
    total_rows = [0]

    if extension == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                parts = archive.infolist()
                if len(parts) > 500 or sum(item.file_size for item in parts) > MAX_UNCOMPRESSED_BYTES:
                    raise SpreadsheetImportError("工作簿展开后过大，请拆分表格再导入")
                if any(item.flag_bits & 1 for item in parts):
                    raise SpreadsheetImportError("不支持加密的 Excel 文件")
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook.worksheets:
                    rows = sheet.iter_rows(min_row=1, max_col=MAX_COLUMNS, values_only=True)
                    _read_table(sheet.title, rows, result, issues, report, default_status, total_rows)
            finally:
                workbook.close()
        except (zipfile.BadZipFile, OSError, ValueError) as error:
            raise SpreadsheetImportError("无法读取该 .xlsx 文件，请用 Excel/WPS 重新保存后再试") from error
    elif extension == ".csv":
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                content = data.decode("gb18030")
            except UnicodeDecodeError as error:
                raise SpreadsheetImportError("CSV 编码无法识别，请另存为 UTF-8 CSV 后再试") from error
        _read_table("CSV数据", csv.reader(io.StringIO(content)), result, issues, report, default_status, total_rows)
    else:
        raise SpreadsheetImportError("仅支持 .xlsx 或 .csv；旧版 .xls 请先另存为 .xlsx")

    if not report and not issues:
        raise SpreadsheetImportError("没有找到手环、无人机、环境或赛段数据表。数据字典不是采集数据表，请按字段填写实际数据行。")
    if report and total_rows[0] == 0:
        raise SpreadsheetImportError("表格只有表头，没有可导入的数据行")
    if issues:
        raise SpreadsheetImportError(f"表格有 {len(issues)} 处问题，本次未导入任何数据", issues[:MAX_ERRORS])
    return PipelineIngestRequest.model_validate(result), report
